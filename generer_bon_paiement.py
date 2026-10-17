# -*- coding: utf-8 -*-
"""Génération Bon de paiement NAP SARL — 1 seule page A4, QR, police professionnelle."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from lxml import etree
from datetime import datetime
from io import BytesIO
import os
import re

try:
    import qrcode
except ImportError:
    qrcode = None

DOSSIER_SORTIE = "bons_paiement"
os.makedirs(DOSSIER_SORTIE, exist_ok=True)

BLEU_FONCE = RGBColor(15, 40, 80)
BLEU_MOYEN = RGBColor(30, 90, 160)
GRIS = RGBColor(55, 55, 55)
GRIS_CLAIR = RGBColor(110, 110, 110)
VERT_MONTANT = RGBColor(34, 139, 34)  # vert pour montant payé
ROUGE_NB = RGBColor(192, 57, 43)  # rouge clair pour NB
BLEU_CIEL = RGBColor(70, 160, 210)  # ligne sous titre


def set_run_font(run, size=9, bold=False, color=None, name="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = name
    r = run._element.rPr
    if r is not None:
        r.rFonts.set(qn("w:eastAsia"), name)
    if color:
        run.font.color.rgb = color


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_margins(cell, top=30, bottom=30, left=50, right=50):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for m, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)



def set_table_borders(table, color="1F4E79", size="8"):
    """Bordures visibles pour un vrai tableau."""
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    if tbl.tblPr is None:
        tbl.insert(0, tblPr)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    # remove old borders if any
    old = tblPr.find(qn("w:tblBorders"))
    if old is not None:
        tblPr.remove(old)
    tblPr.append(borders)


def set_col_widths(table, widths_cm):
    """Largeurs de colonnes en cm."""
    for row in table.rows:
        for idx, w in enumerate(widths_cm):
            if idx < len(row.cells):
                row.cells[idx].width = Cm(w)


def tight_para(p, before=1, after=1):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.0


def generer_qr_png_bytes(texte: str) -> BytesIO:
    if qrcode is None:
        raise ImportError("Installez qrcode : pip install qrcode[pil]")
    qr = qrcode.QRCode(
        version=None,
        error_correction=qrcode.constants.ERROR_CORRECT_M,
        box_size=5,
        border=1,
    )
    qr.add_data(texte)
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white")
    buf = BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    return buf


def ajouter_filigrane_centre(doc, reference):
    date_str = datetime.now().strftime("%d/%m/%Y")
    texte = f"STE NAP SARL  •  {reference}  •  {date_str}"
    section = doc.sections[0]
    header = section.header
    header.is_linked_to_previous = False
    for p in header.paragraphs:
        p.clear()
    paragraph = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    r = OxmlElement("w:r")
    paragraph._p.append(r)
    rPr = OxmlElement("w:rPr")
    rPr.append(OxmlElement("w:noProof"))
    r.append(rPr)
    pict = OxmlElement("w:pict")
    r.append(pict)
    shape_xml = f'''
    <v:shape xmlns:v="urn:schemas-microsoft-com:vml"
             xmlns:o="urn:schemas-microsoft-com:office:office"
             xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
             id="FiligraneNAP"
             type="#_x0000_t136"
             style="position:absolute;margin-left:0;margin-top:0;width:480pt;height:50pt;z-index:-251658240;mso-position-horizontal:center;mso-position-horizontal-relative:page;mso-position-vertical:center;mso-position-vertical-relative:page;"
             fillcolor="#C8C8C8" stroked="f">
      <v:fill opacity=".35"/>
      <v:textbox inset="0,0,0,0">
        <w:txbxContent>
          <w:p>
            <w:pPr><w:jc w:val="center"/></w:pPr>
            <w:r>
              <w:rPr>
                <w:b/><w:sz w:val="24"/><w:szCs w:val="24"/>
                <w:color w:val="A0A0A0"/>
              </w:rPr>
              <w:t xml:space="preserve">{texte}</w:t>
            </w:r>
          </w:p>
        </w:txbxContent>
      </v:textbox>
    </v:shape>
    '''
    pict.append(etree.fromstring(shape_xml))


def _label_val(cell, label, value, label_size=8, val_size=9):
    cell.text = ""
    set_cell_margins(cell)
    p = cell.paragraphs[0]
    tight_para(p, 1, 0)
    r = p.add_run(f"{label}\n")
    set_run_font(r, size=label_size, bold=True, color=BLEU_MOYEN)
    r2 = p.add_run(str(value) if value not in (None, "nan", "None") else "—")
    set_run_font(r2, size=val_size, color=GRIS)


def forcer_une_page(doc):
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.0)
        section.bottom_margin = Cm(1.0)
        section.left_margin = Cm(1.3)
        section.right_margin = Cm(1.3)
        section.header_distance = Cm(0.3)
        section.footer_distance = Cm(0.3)
    for p in doc.paragraphs:
        pf = p.paragraph_format
        pf.space_before = Pt(min(pf.space_before.pt if pf.space_before else 0, 4))
        pf.space_after = Pt(min(pf.space_after.pt if pf.space_after else 0, 3))
        pf.line_spacing = 1.0
        pf.page_break_before = False
        pf.keep_together = True
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(1)
                    p.paragraph_format.line_spacing = 1.0


def generer_bon_paiement(data):
    doc = Document()
    forcer_une_page(doc)

    reference = str(data.get("No_transaction", "")).strip()
    ajouter_filigrane_centre(doc, reference)

    heure_gen = datetime.now().strftime("%H:%M:%S")
    utilisateur_gen = str(
        data.get("Utilisateur") or data.get("Cree_par") or "Systeme"
    ).strip()
    try:
        montant_num = float(data.get("Montant", 0) or 0)
        montant_str = f"{montant_num:,.0f}".replace(",", " ")
    except Exception:
        montant_num = 0
        montant_str = str(data.get("Montant", "0"))
    benef = str(data.get("Beneficiaire") or data.get("Prestataire") or "").strip()

    texte_qr = (
        f"STE NAP SARL|PAY|{reference}|{montant_str}|{benef}|{heure_gen}|{utilisateur_gen}"
    )

    # --- En-tête + QR compact ---
    t_top = doc.add_table(rows=1, cols=2)
    t_top.autofit = True
    c_left, c_right = t_top.rows[0].cells
    shade_cell(c_left, "0F2850")
    set_cell_margins(c_left, top=50, bottom=50, left=80, right=80)
    p = c_left.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tight_para(p, 0, 0)
    r = p.add_run("STE NAP SARL")
    set_run_font(r, size=14, bold=True, color=RGBColor(255, 255, 255))
    p2 = c_left.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tight_para(p2, 0, 0)
    r2 = p2.add_run(
        "PLOT NO 14 GLODJIGBE ECONOMIC ZONE • BENIN\n"
        "RCCM RB/COT/21 B 30710 • IFU 3202113334014 • TEL +229 01 94 94 04 14"
    )
    set_run_font(r2, size=7, color=RGBColor(200, 220, 255))

    c_right.text = ""
    set_cell_margins(c_right, top=20, bottom=20, left=20, right=20)
    pq = c_right.paragraphs[0]
    pq.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if qrcode:
        pq.add_run().add_picture(generer_qr_png_bytes(texte_qr), width=Cm(1.7))
    pl = c_right.add_paragraph()
    pl.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    tight_para(pl, 0, 0)
    set_run_font(pl.add_run("Scan sécurisé NAP"), size=6, color=GRIS_CLAIR)

    # Titre
    titre = doc.add_paragraph()
    titre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tight_para(titre, 4, 2)
    rt = titre.add_run("BON DE PAIEMENT")
    set_run_font(rt, size=13, bold=True, color=BLEU_MOYEN)

    # Ligne bleu ciel sur toute la largeur sous le titre
    ligne = doc.add_paragraph()
    tight_para(ligne, 0, 4)
    pPr = ligne._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "18")  # épaisseur
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "46A0D2")  # bleu ciel
    pBdr.append(bottom)
    pPr.append(pBdr)

    # Infos principales 2x2
    date_v = str(data.get("Date", "")).replace(" 00:00:00", "")
    infos = [
        ("Date", date_v),
        ("Prestataire / Fournisseur", data.get("Prestataire", "")),
        ("N° transaction", reference),
        ("Nom du percepteur", data.get("Beneficiaire", "")),
        ("Type d'opération", data.get("Type_operation", "")),
        ("N° pièce d'identité", data.get("Piece_identite", "")),
        ("Produit", data.get("Produit", "")),
        ("Date d'expiration", data.get("Date_expiration", "")),
        ("", ""),
        ("Adresse / Téléphone", data.get("Telephone", "")),
    ]
    t_info = doc.add_table(rows=5, cols=4)
    t_info.autofit = True
    pairs = [
        (0, 0, infos[0]), (0, 2, infos[1]),
        (1, 0, infos[2]), (1, 2, infos[3]),
        (2, 0, infos[4]), (2, 2, infos[5]),
        (3, 0, infos[6]), (3, 2, infos[7]),
        (4, 0, ("", "")), (4, 2, infos[9]),
    ]
    for row_i, col_i, (lab, val) in pairs:
        if not lab:
            continue
        cell = t_info.rows[row_i].cells[col_i]
        # merge label+value across 2 cols visually by writing in col and next
        cell2 = t_info.rows[row_i].cells[col_i + 1]
        cell.merge(cell2)
        _label_val(cell, lab, val)

    # Détails opérations — vrai tableau bordé (comme la fiche Excel)
    sec = doc.add_paragraph()
    tight_para(sec, 6, 2)
    set_run_font(sec.add_run("Détails des opérations"), size=10, bold=True, color=BLEU_MOYEN)

    # data.get("Lignes") = liste de dicts optionnelle pour plusieurs lignes
    lignes = data.get("Lignes")
    if not lignes:
        lignes = [{
            "Ref_doc": data.get("Ref_doc", "—"),
            "Description": data.get("Description", "—"),
            "Montant": data.get("Montant", 0),
            "Observation": data.get("Observation", "—"),
        }]

    t_det = doc.add_table(rows=1 + len(lignes), cols=4)
    t_det.autofit = False
    set_table_borders(t_det, color="1F4E79", size="8")
    set_col_widths(t_det, [4.2, 7.5, 3.5, 2.8])

    headers = ["Réf. Doc", "Description / Objet du paiement", "Montant", "Observation"]
    for i, h in enumerate(headers):
        cell = t_det.rows[0].cells[i]
        shade_cell(cell, "1F4E79")
        set_cell_margins(cell, top=30, bottom=30, left=40, right=40)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        tight_para(p, 0, 0)
        set_run_font(p.add_run(h), size=8, bold=True, color=RGBColor(255, 255, 255))

    for r_idx, ligne in enumerate(lignes, start=1):
        try:
            m = float(ligne.get("Montant", 0) or 0)
            m_txt = f"XOF {m:,.0f}".replace(",", " ")
        except Exception:
            m_txt = str(ligne.get("Montant", "—"))
        vals = [
            ligne.get("Ref_doc", "—"),
            ligne.get("Description", "—"),
            m_txt,
            ligne.get("Observation", "—"),
        ]
        # alternance léger gris
        fill = "F5F8FC" if r_idx % 2 == 0 else "FFFFFF"
        for i, v in enumerate(vals):
            cell = t_det.rows[r_idx].cells[i]
            shade_cell(cell, fill)
            set_cell_margins(cell, top=35, bottom=35, left=40, right=40)
            p = cell.paragraphs[0]
            tight_para(p, 1, 1)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i >= 2 else WD_ALIGN_PARAGRAPH.LEFT
            txt = str(v) if v not in (None, "nan", "None", "") else "—"
            set_run_font(p.add_run(txt), size=9, bold=(i == 2), color=(VERT_MONTANT if i == 2 else GRIS))

    # Total — montant en vert clair
    tot = doc.add_paragraph()
    tight_para(tot, 4, 1)
    tot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(tot.add_run("Montant total : "), size=9, bold=True, color=BLEU_FONCE)
    set_run_font(tot.add_run(f"XOF {montant_str}"), size=11, bold=True, color=VERT_MONTANT)

    lettres = str(data.get("Montant_lettres") or "").strip()
    if lettres:
        pl = doc.add_paragraph()
        tight_para(pl, 1, 2)
        set_run_font(pl.add_run("Montant en lettres : "), size=8, bold=True, color=GRIS_CLAIR)
        set_run_font(pl.add_run(lettres), size=9, bold=True, color=GRIS)

    # Mode paiement
    t_pay = doc.add_table(rows=1, cols=3)
    t_pay.autofit = True
    for i, (lab, key) in enumerate([
        ("Mode paiement", "Mode_paiement"),
        ("N° chèque", "Numero_cheque"),
        ("Banque", "Banque"),
    ]):
        _label_val(t_pay.rows[0].cells[i], lab, data.get(key, "—"))

    # Signatures
    sig = doc.add_paragraph()
    tight_para(sig, 8, 2)
    set_run_font(sig.add_run("Signatures"), size=9, bold=True, color=BLEU_MOYEN)

    t_sig = doc.add_table(rows=2, cols=3)
    t_sig.autofit = True
    for i, lab in enumerate(["Manager logistics", "Comptabilité", "Bénéficiaire"]):
        cell = t_sig.rows[0].cells[i]
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        tight_para(p, 1, 1)
        set_run_font(p.add_run(lab), size=8, bold=True, color=BLEU_FONCE)
        cell2 = t_sig.rows[1].cells[i]
        p2 = cell2.paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        tight_para(p2, 1, 1)
        set_run_font(p2.add_run("\n\nvise / date\n_______________"), size=8, color=GRIS_CLAIR)

    # Note légale — NB en rouge clair (sans mention "Document généré…")
    note = doc.add_paragraph()
    tight_para(note, 6, 1)
    set_run_font(
        note.add_run(
            "NB : Toute remise de chèque doit être confirmée par signature du bénéficiaire "
            "et accompagnée d'une copie du chèque et pièce pour fins de traçabilité."
        ),
        size=8,
        bold=True,
        color=ROUGE_NB,
    )

    forcer_une_page(doc)

    safe_ref = re.sub(r'[\\/*?:"<>|]', "", reference) or "PAIEMENT"
    nom_fichier = f"Bon_Paiement_{safe_ref}.docx"
    chemin = os.path.join(DOSSIER_SORTIE, nom_fichier)
    doc.save(chemin)
    return chemin


if __name__ == "__main__":
    demo = {
        "Date": "15/10/2026",
        "No_transaction": "CHQ-0389-2026",
        "Type_operation": "PURCHASE",
        "Produit": "SHEA",
        "Prestataire": "AZOGA ARMAND",
        "Beneficiaire": "AZOGA ARMAND",
        "Piece_identite": "100404882",
        "Date_expiration": "05/04/2028",
        "Telephone": "-",
        "Ref_doc": "CC2315RB_17.86 MT",
        "Description": "PURCHASE OF SHEA",
        "Montant": 5843800,
        "Observation": "PAYE",
        "Montant_lettres": "Cinq millions huit cent quarante-trois mille huit cents francs CFA",
        "Mode_paiement": "CHEQUE",
        "Numero_cheque": "2000751",
        "Banque": "ECOBANK BENIN",
        "Utilisateur": "Admin",
    }
    print(generer_bon_paiement(demo))
