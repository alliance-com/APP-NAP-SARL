"""
moteur.py - Cœur du rapprochement bancaire NAP SARL (aucune dépendance Streamlit).

  lire_releve_pdf(fichier)   -> (meta, DataFrame des lignes de la banque)
  lire_excel(fichier)        -> DataFrame des lignes du classeur interne (feuille STATEMENT)
  rapprocher(...)            -> DataFrames rapprochés + synthèse chiffrée
"""
import io
import re
import unicodedata
from datetime import datetime, date, timedelta

import pandas as pd
import pdfplumber
from openpyxl import load_workbook

# --------------------------------------------------------------------------
# Outils
# --------------------------------------------------------------------------
RE_DATE = re.compile(r"^\d{2}/\d{2}/\d{4}$")
RE_CHEQUE = re.compile(r"(?:CHQ|CHEQUE|CHEQ)\s*N\W{0,3}\s*(\d{5,8})", re.I)

def _entier(txt):
    chiffres = re.sub(r"[^\d]", "", str(txt))
    return int(chiffres) if chiffres else 0

def _sans_accent(s):
    s = unicodedata.normalize("NFKD", str(s))
    return "".join(c for c in s if not unicodedata.combining(c))

def _jetons(s):
    """Mots significatifs (>=3 lettres) en majuscules, sans accents, pour comparer des noms."""
    s = _sans_accent(s).upper()
    s = re.sub(r"[^A-Z ]", " ", s)
    stop = {"RETRAIT", "PAR", "CHEQUE", "CHQ", "FAV", "PMT", "PAID", "SELF", "CNI", "CIP", "THE", "DES",
            "PAIEMENT", "NAP", "SARL", "EBJ", "AGENCE", "REMISE", "ECOBANK", "ENVOYE", "ENCAISSEMENT",
            "PURCHASE", "SHEA", "SOYA", "VOTRE", "TRANSFERT", "FONDS"}
    return {m for m in s.split() if len(m) >= 3 and m not in stop}

def similarite_noms(a, b):
    ja, jb = _jetons(a), _jetons(b)
    if not ja or not jb:
        return 0.0
    return len(ja & jb) / min(len(ja), len(jb))

def fmt_xof(n):
    if n is None or (isinstance(n, float) and pd.isna(n)):
        return ""
    return f"{int(round(n)):,}".replace(",", " ")

# --------------------------------------------------------------------------
# 1) Relevé PDF Ecobank
# --------------------------------------------------------------------------
def _nettoyer_libelle(t):
    t = t.replace("Â°", "°").replace("Â", "")
    t = re.sub(r"//\s*([A-Z])\s+([A-Z]{3,})", lambda m: "// " + m.group(1) + m.group(2), t)
    t = re.sub(r"\s+", " ", t).strip()
    return t

def _extraire_beneficiaire(narr):
    if not re.search(r"CHQ|CHEQUE", narr, re.I):
        return ""
    m = re.search(r"\b(?:FAV|FV)\b[\s/]*(.+)", narr, re.I)
    if not m:
        return ""
    txt = m.group(1)
    txt = re.split(r"\b(?:CNI|CIP|CB|PP|PAID|BD)\b|\d{6,}", txt, maxsplit=1, flags=re.I)[0]
    txt = re.sub(r"[/]+", " ", txt)
    return re.sub(r"\s+", " ", txt).strip(" -:")

def _classer(narr, debit, credit):
    n = _sans_accent(narr).upper()
    if ("COMMISSION" in n or "FRAIS" in n) or ("TAXE" in n and "VIREMENT SALAIRE" in n):
        return "FRAIS BANCAIRE", "Frais bancaires"
    if "VIREMENT SALAIRE" in n:
        return "SALAIRES", "Virement salaires"
    if "DGI" in n:
        return "IMPOTS / DGI", "Paiement DGI"
    if "TIRAGE" in n and ("CREDIT" in n or "CCT" in n):
        return "FINANCEMENT", "Tirage ligne de crédit"
    if "BILL PAYMENTS" in n or "COMPTES D ATTENTE" in n:
        return "FACTURES (OMNIS)", "Paiement facture / Omnis"
    if "TRANSFERT DE FONDS - RECUS" in n or "TRANSFERT DE FONDS-RECUS" in n or (credit and "TRF" in n):
        m = re.search(r"B/O\s*(?:\d/)?\s*(.+?)(?:\s+IRO|\s+INVOICE|\s+IFO|$)", narr, re.I)
        return "VIREMENT RECU", ("Virement reçu - " + m.group(1).strip()) if m else "Virement reçu"
    if "ENCAISSEMENT" in n or "REMISE CHQ" in n or "REMISE DE CHEQUE" in n or "REMISE CHEQUE" in n:
        return "REMISE CHEQUE", "Remise de chèque"
    if "RETRAIT PAR CHEQUE" in n:
        return "CHEQUE PAYE", "Chèque payé"
    if re.match(r"^\d{5}IW\d+//\d+//$", narr.strip()):
        return "A IDENTIFIER", "Opération sans libellé (à identifier)"
    return "AUTRE", "Autre opération"

def lire_releve_pdf(fichier):
    """fichier = chemin ou objet fichier (upload Streamlit). Retourne (meta, DataFrame)."""
    meta, lignes = {}, []
    with pdfplumber.open(fichier) as pdf:
        p0 = pdf.pages[0].extract_text() or ""

        def champ(motif):
            m = re.search(motif, p0)
            return m.group(1).strip() if m else ""

        meta["societe"] = champ(r"Company Name\s*:?\s*(.+?)(?:\s{2,}|\n|Financial)") or "STE NAP SARL"
        meta["banque"] = champ(r"Financial Institution\s*:?\s*(Ecobank\s*\w+)") or "Ecobank Benin"
        meta["compte"] = champ(r"Account\s*:\s*(\d{8,})")
        meta["nom_compte"] = champ(r"Account Name\s*:?\s*(.+?)(?:\s+Closing|\n)")
        meta["devise"] = champ(r"Account Currency\s*:?\s*(\w+)") or "XOF"
        meta["solde_ouverture"] = _entier(champ(r"Opening Book Balance\s*:?\s*XOF\s*([\d,]+)"))
        meta["solde_cloture"] = _entier(champ(r"Closing Book Balance\s*:?\s*XOF\s*([\d,]+)"))
        meta["total_retraits"] = _entier(champ(r"Total Withdrawals\s*:?\s*XOF\s*([\d,]+)"))
        meta["total_depots"] = _entier(champ(r"Total Deposit\s*:?\s*XOF\s*([\d,]+)"))
        m = re.search(r"BETWEEN\s+(\d{2}/\d{2}/\d{4})\s+to\s+(\d{2}/\d{2}/\d{4})", p0)
        if m:
            meta["periode_debut"] = datetime.strptime(m.group(1), "%m/%d/%Y").date()
            meta["periode_fin"] = datetime.strptime(m.group(2), "%m/%d/%Y").date()

        for page in pdf.pages:
            mots = page.extract_words()
            entete = [w for w in mots if w["text"] == "S.No."]
            if not entete:
                continue
            haut = entete[0]["top"] + 10
            bas = page.height - 45
            for w in mots:
                if w["text"] in ("---End",) and w["top"] > haut:
                    bas = min(bas, w["top"] - 2)
            corps = [w for w in mots if haut < w["top"] < bas]
            debuts = sorted([w for w in corps if w["x0"] < 75 and w["text"].isdigit()], key=lambda w: w["top"])
            for i, d in enumerate(debuts):
                y0 = d["top"] - 2
                y1 = debuts[i + 1]["top"] - 2 if i + 1 < len(debuts) else bas
                seg = [w for w in corps if y0 <= w["top"] < y1]
                tri = sorted(seg, key=lambda w: (round(w["top"]), w["x0"]))
                narr = " ".join(w["text"] for w in tri if 165 <= w["x0"] < 320)
                dates_post = [w["text"] for w in tri if w["x0"] < 165 and RE_DATE.match(w["text"])]
                dates_val = [w["text"] for w in tri if 320 <= w["x0"] < 440 and RE_DATE.match(w["text"])]
                debit = credit = solde = 0
                for w in seg:
                    if w["text"].isdigit() and w["x0"] >= 430:
                        if w["x1"] < 540:
                            debit = int(w["text"])
                        elif w["x1"] < 690:
                            credit = int(w["text"])
                        else:
                            solde = int(w["text"])
                narr = _nettoyer_libelle(narr)
                cat, lib = _classer(narr, debit, credit)
                lignes.append({
                    "No": int(d["text"]),
                    "Date": datetime.strptime(dates_post[0], "%d/%m/%Y").date() if dates_post else None,
                    "Date valeur": datetime.strptime(dates_val[0], "%d/%m/%Y").date() if dates_val else None,
                    "Libellé banque": narr,
                    "Débit": debit,
                    "Crédit": credit,
                    "Solde banque": solde,
                    "N° chèque": (RE_CHEQUE.search(narr).group(1) if RE_CHEQUE.search(narr) else ""),
                    "Bénéficiaire": _extraire_beneficiaire(narr),
                    "Catégorie suggérée": cat,
                    "Libellé propre": lib,
                })
    df = pd.DataFrame(lignes).sort_values("No").reset_index(drop=True)
    return meta, df

def controler_releve(meta, df):
    """Vérifie que le relevé PDF a été lu sans erreur. Retourne une liste de (ok, message)."""
    res = []
    ouv = meta.get("solde_ouverture", 0)
    courant, fautes = ouv, []
    for _, r in df.iterrows():
        courant = courant - r["Débit"] + r["Crédit"]
        if courant != r["Solde banque"]:
            fautes.append(int(r["No"]))
            courant = r["Solde banque"]
    res.append((not fautes, "Chaîne des soldes ligne par ligne cohérente" if not fautes
                else f"Solde incohérent aux lignes n° {fautes[:10]}"))
    ok_d = df["Débit"].sum() == meta.get("total_retraits")
    ok_c = df["Crédit"].sum() == meta.get("total_depots")
    res.append((ok_d, f"Total débits lus = {fmt_xof(df['Débit'].sum())} (relevé : {fmt_xof(meta.get('total_retraits'))})"))
    res.append((ok_c, f"Total crédits lus = {fmt_xof(df['Crédit'].sum())} (relevé : {fmt_xof(meta.get('total_depots'))})"))
    final = ouv + df["Crédit"].sum() - df["Débit"].sum()
    res.append((final == meta.get("solde_cloture"),
                f"Ouverture + crédits - débits = {fmt_xof(final)} (solde de clôture : {fmt_xof(meta.get('solde_cloture'))})"))
    return res

# --------------------------------------------------------------------------
# 2) Classeur Excel interne (feuille STATEMENT)
# --------------------------------------------------------------------------
def _en_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        for f in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
            try:
                return datetime.strptime(v.strip(), f).date()
            except ValueError:
                pass
    return None

def _en_montant(v):
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    t = str(v).strip()
    if t.startswith("="):
        try:
            if re.fullmatch(r"=[\d\.\*\+\-\/\(\) ]+", t):
                return float(eval(t[1:], {"__builtins__": {}}, {}))
        except Exception:
            return 0.0
        return 0.0
    try:
        return float(re.sub(r"[^\d\.\-]", "", t.replace(",", ".")))
    except ValueError:
        return 0.0

def lire_excel(fichier, feuille="STATEMENT"):
    """Lit la feuille STATEMENT (colonnes A à H, en-têtes en ligne 3)."""
    wb = load_workbook(fichier, data_only=True)
    if feuille not in wb.sheetnames:
        raise ValueError(f"Feuille « {feuille} » introuvable. Feuilles : {', '.join(wb.sheetnames)}")
    ws = wb[feuille]
    ligne_entete = next((r for r in range(1, 15)
                         if any(str(ws.cell(r, c).value or "").strip().upper() == "DESCRIPTION" for c in range(1, 9))), 3)
    lignes = []
    for r, row in enumerate(ws.iter_rows(min_row=ligne_entete + 1, max_col=8, values_only=True), start=ligne_entete + 1):
        d = _en_date(row[0])
        desc = (row[1] or "")
        if d is None and not str(desc).strip():
            continue
        cheque = row[5]
        if isinstance(cheque, float) and cheque.is_integer():
            cheque = int(cheque)
        lignes.append({
            "Ligne Excel": r,
            "Date": d,
            "Description": str(desc).strip(),
            "Crédit": _en_montant(row[2]),
            "Débit": _en_montant(row[3]),
            "Solde Excel": _en_montant(row[4]),
            "N° chèque": str(cheque).strip() if cheque not in (None, "") else "",
            "Motif": str(row[6] or "").strip(),
            "Catégorie": str(row[7] or "").strip(),
        })
    return pd.DataFrame(lignes)

# --------------------------------------------------------------------------
# 3) Rapprochement
# --------------------------------------------------------------------------
STATUT_OK = "Rapproché"
STATUT_VERIF = "À vérifier"
STATUT_ECART = "Écart de montant"
STATUT_ABSENT = "Absent d'Excel"
STATUT_EXCEL_SEUL = "Pas encore en banque"

def _cheques_compatibles(c1, c2):
    """Deux n° de chèque sont compatibles s'ils sont égaux, absents, ou si l'un est tronqué."""
    if not c1 or not c2 or c1 == c2:
        return True
    return c1.startswith(c2) or c2.startswith(c1)

def rapprocher(meta, banque, excel, date_depart, solde_depart, tolerance_jours=7, jours_avant=45):
    """
    banque, excel : DataFrames issus des fonctions ci-dessus.
    date_depart   : début de la période contrôlée (incluse).
    solde_depart  : solde de clôture du contrôle précédent (point de départ commun).
    Règles (dans l'ordre, une ligne banque = une ligne Excel) :
      1. même n° de chèque + même montant                       -> Rapproché
      2. même n° de chèque, montant différent                   -> Écart de montant
      3. même montant + même sens + nom proche + dates proches  -> Rapproché
      4. même montant + même sens + dates proches (sans nom)    -> Rapproché, ou À vérifier si douteux
      5. nom proche + montant à moins de 2 % (change, frais)    -> Écart de montant
    """
    date_fin = meta.get("periode_fin") or banque["Date"].max()
    b = banque.copy()
    b["Montant"] = b["Crédit"] - b["Débit"]
    b["Statut"] = STATUT_ABSENT
    b["Méthode"] = ""
    b["Ligne Excel liée"] = pd.NA
    b["Détail"] = ""
    b["Écart"] = 0.0

    e = excel.copy()
    e["Montant"] = e["Crédit"] - e["Débit"]
    e["En période"] = e["Date"].apply(lambda d: d is not None and date_depart <= d <= date_fin)
    debut_pool = date_depart - timedelta(days=jours_avant)
    e["Candidat"] = e["Date"].apply(lambda d: d is not None and debut_pool <= d <= date_fin + timedelta(days=tolerance_jours))
    e["Statut"] = STATUT_EXCEL_SEUL
    e["Ligne banque liée"] = pd.NA

    libres_e = set(e.index[e["Candidat"]])
    libres_b = set(b.index)

    def lier(ib, ie, statut, methode, detail=""):
        b.at[ib, "Statut"] = statut
        b.at[ib, "Méthode"] = methode
        b.at[ib, "Ligne Excel liée"] = int(e.at[ie, "Ligne Excel"])
        b.at[ib, "Détail"] = detail
        e.at[ie, "Statut"] = "Rapproché"
        e.at[ie, "Ligne banque liée"] = int(b.at[ib, "No"])
        libres_e.discard(ie)
        libres_b.discard(ib)

    def jours(ib, ie):
        d1, d2 = b.at[ib, "Date"], e.at[ie, "Date"]
        return abs((d1 - d2).days) if d1 and d2 else 999

    def texte_b(ib):
        return b.at[ib, "Bénéficiaire"] + " " + b.at[ib, "Libellé banque"]

    # Règle 1 : même chèque + même montant
    for ib in sorted(libres_b):
        chq = b.at[ib, "N° chèque"]
        if not chq:
            continue
        for ie in sorted(libres_e):
            if e.at[ie, "N° chèque"] == chq and abs(e.at[ie, "Montant"] - b.at[ib, "Montant"]) < 0.5:
                lier(ib, ie, STATUT_OK, "N° chèque + montant")
                break

    # Règle 2 : même chèque, montant différent
    for ib in sorted(libres_b):
        chq = b.at[ib, "N° chèque"]
        if not chq:
            continue
        for ie in sorted(libres_e):
            if e.at[ie, "N° chèque"] == chq:
                diff = b.at[ib, "Montant"] - e.at[ie, "Montant"]
                lier(ib, ie, STATUT_ECART, "N° chèque",
                     f"Chèque {chq} : banque {fmt_xof(abs(b.at[ib, 'Montant']))} / Excel {fmt_xof(abs(e.at[ie, 'Montant']))}")
                b.at[ib, "Écart"] = diff
                break

    # Règles 3 et 4 : même montant et même sens
    for avec_nom in (True, False):
        for ib in sorted(libres_b):
            candidats = []
            for ie in libres_e:
                if abs(e.at[ie, "Montant"] - b.at[ib, "Montant"]) >= 0.5:
                    continue
                sim = similarite_noms(e.at[ie, "Description"], texte_b(ib))
                dj = jours(ib, ie)
                chq_ok = _cheques_compatibles(b.at[ib, "N° chèque"], e.at[ie, "N° chèque"])
                if not chq_ok and sim < 0.5:
                    continue
                if dj > tolerance_jours and sim < 0.5:
                    continue
                if avec_nom and sim < 0.5:
                    continue
                candidats.append((-sim, dj, ie, sim, chq_ok))
            if not candidats:
                continue
            candidats.sort()
            _, dj, ie, sim, chq_ok = candidats[0]
            raisons = []
            if sim < 0.5 and len(candidats) > 1:
                raisons.append(f"{len(candidats)} écritures Excel de même montant")
            if b.at[ib, "Catégorie suggérée"] == "A IDENTIFIER":
                raisons.append("libellé banque vide")
            if (b.at[ib, "N° chèque"] and e.at[ie, "N° chèque"] and b.at[ib, "N° chèque"] != e.at[ie, "N° chèque"]):
                raisons.append(f"n° chèque banque {b.at[ib, 'N° chèque']} / Excel {e.at[ie, 'N° chèque']}")
            if dj > tolerance_jours:
                raisons.append(f"{dj} jours d'écart")
            methode = "Montant + nom" if sim >= 0.5 else "Montant seul"
            lier(ib, ie, STATUT_VERIF if raisons else STATUT_OK, methode, " ; ".join(raisons))

    # Règle 5 : nom proche + montant presque égal (change, frais) -> écart de montant
    for ib in sorted(libres_b):
        for ie in sorted(libres_e):
            mb, me = b.at[ib, "Montant"], e.at[ie, "Montant"]
            if mb == 0 or me == 0 or (mb > 0) != (me > 0):
                continue
            if abs(mb - me) / max(abs(mb), abs(me)) > 0.02:
                continue
            if similarite_noms(e.at[ie, "Description"], texte_b(ib)) < 0.5 or jours(ib, ie) > tolerance_jours:
                continue
            lier(ib, ie, STATUT_ECART, "Nom + montant proche",
                 f"Banque {fmt_xof(abs(mb))} / Excel {fmt_xof(abs(me))} (change ou frais ?)")
            b.at[ib, "Écart"] = mb - me
            break

    absentes = b["Statut"] == STATUT_ABSENT
    b.loc[absentes, "Détail"] = "À saisir dans Excel"
    b.loc[absentes & (b["Catégorie suggérée"] == "A IDENTIFIER"), "Détail"] = "Libellé banque vide : à identifier, puis saisir dans Excel"

    # ---------------- Chiffres de synthèse ----------------
    b_per = b[b["Date"].apply(lambda d: d is not None and date_depart <= d <= date_fin)]
    e_per = e[e["En période"]]

    credits_banque = int(b_per["Crédit"].sum())
    debits_banque = int(b_per["Débit"].sum())
    credits_excel = float(e_per["Crédit"].sum())
    debits_excel = float(e_per["Débit"].sum())

    solde_banque_cloture = meta.get("solde_cloture", 0)
    solde_excel_cloture = solde_depart + credits_excel - debits_excel
    ecart_total = solde_banque_cloture - solde_excel_cloture

    banque_seule = b_per[b_per["Statut"] == STATUT_ABSENT]
    lies = b_per[b_per["Statut"] != STATUT_ABSENT]
    b_lies_hors = lies[~lies["Ligne Excel liée"].isin(e_per["Ligne Excel"])]
    excel_seul = e_per[e_per["Statut"] == STATUT_EXCEL_SEUL]
    ecarts_montant = b_per[b_per["Statut"] == STATUT_ECART]

    net_banque_seule = int(banque_seule["Montant"].sum())
    net_hors_periode = int(b_lies_hors["Montant"].sum())
    net_excel_seul = float(excel_seul["Montant"].sum())
    net_ecarts_montant = float(ecarts_montant["Écart"].sum())
    ecart_ouverture = meta.get("solde_ouverture", 0) - solde_depart

    ecart_explique = ecart_ouverture + net_banque_seule + net_hors_periode - net_excel_seul + net_ecarts_montant
    non_explique = round(ecart_total - ecart_explique, 2)

    dernier_solde = e[e["Date"].apply(lambda d: d is not None and d <= date_fin)]["Solde Excel"]
    synthese = {
        "date_depart": date_depart, "date_fin": date_fin,
        "solde_depart": solde_depart, "solde_ouverture_banque": meta.get("solde_ouverture", 0),
        "credits_banque": credits_banque, "debits_banque": debits_banque,
        "solde_banque_cloture": solde_banque_cloture,
        "credits_excel": round(credits_excel, 2), "debits_excel": round(debits_excel, 2),
        "solde_excel_cloture": round(solde_excel_cloture, 2),
        "solde_excel_colonne": float(dernier_solde.iloc[-1]) if len(dernier_solde) else 0.0,
        "ecart_total": round(ecart_total, 2),
        "ecart_ouverture": ecart_ouverture,
        "net_banque_seule": net_banque_seule, "nb_banque_seule": len(banque_seule),
        "net_hors_periode": net_hors_periode, "nb_hors_periode": len(b_lies_hors),
        "net_excel_seul": round(net_excel_seul, 2), "nb_excel_seul": len(excel_seul),
        "net_ecarts_montant": round(net_ecarts_montant, 2), "nb_ecarts_montant": len(ecarts_montant),
        "ecart_explique": round(ecart_explique, 2), "non_explique": non_explique,
        "nb_lignes_banque": len(b_per), "nb_rapprochees": int((b_per["Statut"] == STATUT_OK).sum()),
        "nb_a_verifier": int((b_per["Statut"] == STATUT_VERIF).sum()),
    }
    return {"banque": b, "excel": e, "synthese": synthese}